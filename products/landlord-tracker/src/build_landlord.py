"""Build the UK Landlord Rental Tracker workbooks (blank + example).

Reuses the Sole Trader tracker's helpers (same Settings quarter layout, so the
shared quarter formula works unchanged). Plain formulas only.
"""
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "sole-trader-tracker" / "src"))
import build_xlsx as st  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

FIRST, LAST = st.FIRST, st.LAST
font, header, title, unlock, BOX, GBP, DMY = st.font, st.header, st.title, st.unlock, st.BOX, st.GBP, st.DMY
INPUT, BRAND, BRAND_SOFT = st.INPUT, st.BRAND, st.BRAND_SOFT
QUARTERS = st.QUARTERS

FINANCE = "Residential property finance costs"
CATEGORIES = [  # SA105 UK property expense lines, finance costs last (held apart)
    "Rent, rates, insurance and ground rents",
    "Property repairs and maintenance",
    "Non-residential property finance costs",
    "Legal, management and other professional fees",
    "Costs of services provided, including wages",
    "Costs of replacing domestic items",
    "Other allowable property expenses",
    FINANCE,
]
DEDUCTIBLE = CATEGORIES[:-1]
INCOME_TYPES = ["Rent", "Other property income"]
PROPS = 5
PROP_FIRST = 20  # Settings rows 20-24 hold the properties


def build_settings(wb):
    ws = wb.create_sheet("Settings")
    title(ws, "Settings", "Fill in the cream cells once. Everything else updates itself.")
    rows = [
        ("Your name", "Your name", None),
        ("Tax year starts (year)", st.TAX_YEAR, "Type 2026 for the 2026-27 tax year (6 April 2026 to 5 April 2027)."),
        ("Quarter dates", "Standard (6th to 5th)", "Choose Calendar only if you elected calendar quarters with HMRC."),
        ("Self-employment turnover this year £", 0, "Optional. Used only for the Making Tax Digital check on Year End."),
    ]
    for i, (label, val, note) in enumerate(rows):
        r = 4 + i
        ws.cell(row=r, column=2, value=label).font = font(10, True)
        c = ws.cell(row=r, column=3, value=val)
        c.fill = PatternFill("solid", fgColor=INPUT)
        c.border = BOX
        unlock(c)
        if i == 3:
            c.number_format = GBP
        if note:
            ws.cell(row=r, column=4, value=note).font = font(9, color=st.MUTED, italic=True)
    ws["B9"], ws["C9"] = "Tax year", '=C5&"-"&RIGHT(C5+1,2)'
    ws["B10"], ws["C10"] = "Quarter start day", '=IF(C6="Calendar (1st to end)",1,6)'
    header(ws, 12, ["Quarter", "Starts", "Ends", "Update due by"])
    q = [(4, 0, 7, 0, 8, 0), (7, 0, 10, 0, 11, 0), (10, 0, 1, 1, 2, 1), (1, 1, 4, 1, 5, 1)]
    for i, (sm, sy, em, ey, dm, dy) in enumerate(q):
        r = 13 + i
        ws.cell(row=r, column=2, value=QUARTERS[i]).font = font(10, True)
        ws.cell(row=r, column=3, value=f"=DATE($C$5+{sy},{sm},$C$10)").number_format = DMY
        ws.cell(row=r, column=4, value=f"=DATE($C$5+{ey},{em},$C$10)-1").number_format = DMY
        ws.cell(row=r, column=5, value=f"=DATE($C$5+{dy},{dm},7)").number_format = DMY
        for col in range(2, 6):
            ws.cell(row=r, column=col).border = BOX
    ws["B17"] = "Final declaration (year end) due by"
    ws["E17"] = "=DATE($C$5+2,1,31)"
    ws["E17"].number_format = DMY

    header(ws, 19, ["Property (short name)", "Your share %", "Notes"])
    for i in range(PROPS):
        r = PROP_FIRST + i
        for col in (2, 3, 4):
            c = ws.cell(row=r, column=col)
            c.fill = PatternFill("solid", fgColor=INPUT)
            c.border = BOX
            unlock(c)
        ws.cell(row=r, column=3).number_format = "0%"
    ws.cell(row=25, column=2, value="Jointly owned? Enter your share. Spouses and civil partners are taxed 50:50 "
            "unless they have sent HMRC Form 17.").font = font(9, color=st.MUTED, italic=True)
    ws["B27"], ws["C27"] = "Property allowance £", 1000
    ws["C27"].fill = PatternFill("solid", fgColor=INPUT)
    ws["C27"].number_format = GBP
    unlock(ws["C27"])

    ws["G3"] = "Dropdown lists"
    ws["G3"].font = font(10, True, st.MUTED)
    for i, v in enumerate(CATEGORIES):
        ws.cell(row=4 + i, column=7, value=v).font = font(9, color=st.MUTED)
    for i, v in enumerate(INCOME_TYPES):
        ws.cell(row=4 + i, column=9, value=v).font = font(9, color=st.MUTED)
    for i, v in enumerate(["Standard (6th to 5th)", "Calendar (1st to end)"]):
        ws.cell(row=4 + i, column=12, value=v).font = font(9, color=st.MUTED)
    dv = DataValidation(type="list", formula1="=$L$4:$L$5", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C6")
    for col, w in {"B": 40, "C": 24, "D": 30, "E": 16, "F": 3, "G": 46, "H": 2, "I": 22, "L": 22}.items():
        ws.column_dimensions[col].width = w
    ws.protection.sheet = True
    return ws


def share_column(ws):
    """Column I: your share of each row = amount x your % for that property."""
    ws["I4"] = "Your share £"
    ws["I4"].font = font(10, True, "FFFFFF")
    ws["I4"].fill = PatternFill("solid", fgColor=BRAND)
    ws["I4"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions["I"].width = 14
    for r in range(FIRST, LAST + 1):
        c = ws[f"I{r}"]
        c.value = f'=IF(F{r}="","",F{r}*IFERROR(VLOOKUP(C{r},Settings!$B${PROP_FIRST}:$C${PROP_FIRST + PROPS - 1},2,FALSE),0))'
        c.number_format = GBP
        c.font = font(10, color=st.MUTED)
        c.border = BOX


def summary(wb):
    ws = wb.create_sheet("Quarterly Summary")
    title(ws, "Quarterly Summary", "Your share of every property, for each quarterly update. Nothing to type here.")
    ws["B3"] = '="Tax year "&Settings!C9&"   ·   "&Settings!C4'
    ws["B3"].font = font(10, True, st.MUTED)
    header(ws, 5, ["Your share", "Q1", "Q2", "Q3", "Q4", "Year total", "", "To end of Q1", "To end of Q2", "To end of Q3", "To end of Q4"])
    ws["H5"].fill = PatternFill(fill_type=None)
    ws["I4"] = "Year to date (cumulative, as quarterly updates use)"
    ws["I4"].font = font(9, True, st.MUTED)
    ws["B6"] = "Update due by"
    ws["B6"].font = font(10, True, st.WARN)
    for i in range(4):
        c = ws.cell(row=6, column=3 + i, value=f"=Settings!E{13 + i}")
        c.number_format = DMY
        c.font = font(10, True, st.WARN)
    plan = [("INCOME", "sec")] + [(t, "inc") for t in INCOME_TYPES] + [("Total income", "ti"), ("ALLOWABLE EXPENSES", "sec")]
    plan += [(c, "exp") for c in DEDUCTIBLE] + [("Total allowable expenses", "te"), ("NET PROPERTY PROFIT", "np"),
                                                (FINANCE + " (not deducted)", "fin")]
    pos, r = {}, 7
    for label, kind in plan:
        ws.cell(row=r, column=2, value=label)
        pos.setdefault(kind, []).append(r)
        if kind == "sec":
            ws.cell(row=r, column=2).font = font(9, True, BRAND)
        r += 1
    last = r - 1
    for rr in range(7, last + 1):
        kind = next(k for k, v in pos.items() if rr in v)
        if kind == "sec":
            continue
        for qi, q in enumerate(QUARTERS):
            L = "CDEF"[qi]
            if kind == "inc":
                f = f'=SUMIFS(Income!$I${FIRST}:$I${LAST},Income!$E${FIRST}:$E${LAST},$B{rr},Income!$G${FIRST}:$G${LAST},"{q}")'
            elif kind == "exp":
                f = f'=SUMIFS(Expenses!$I${FIRST}:$I${LAST},Expenses!$E${FIRST}:$E${LAST},$B{rr},Expenses!$G${FIRST}:$G${LAST},"{q}")'
            elif kind == "fin":
                f = f'=SUMIFS(Expenses!$I${FIRST}:$I${LAST},Expenses!$E${FIRST}:$E${LAST},"{FINANCE}",Expenses!$G${FIRST}:$G${LAST},"{q}")'
            elif kind == "ti":
                f = f"=SUM({L}{pos['inc'][0]}:{L}{pos['inc'][-1]})"
            elif kind == "te":
                f = f"=SUM({L}{pos['exp'][0]}:{L}{pos['exp'][-1]})"
            else:
                f = f"={L}{pos['ti'][0]}-{L}{pos['te'][0]}"
            ws.cell(row=rr, column=3 + qi, value=f)
        ws.cell(row=rr, column=7, value=f"=SUM(C{rr}:F{rr})")
        for k in range(4):
            ws.cell(row=rr, column=9 + k, value=f"=SUM($C{rr}:{'CDEF'[k]}{rr})")
        for col in list(range(2, 8)) + list(range(9, 13)):
            c = ws.cell(row=rr, column=col)
            c.border = BOX
            if col > 2:
                c.number_format = GBP
            if kind in ("ti", "te", "np"):
                c.font = font(10, True)
                c.fill = PatternFill("solid", fgColor=BRAND_SOFT)
            if kind == "fin":
                c.font = font(9, color=st.MUTED, italic=True)
    notes = ["Residential finance costs (mortgage interest) are not deducted from profit. They give a basic-rate (20%) tax "
             "reduction instead, so they are kept apart here.",
             "Send these figures with your HMRC-recognised bridging or MTD software. This spreadsheet does not submit anything to HMRC."]
    for i, n in enumerate(notes):
        ws.cell(row=last + 2 + i, column=2, value=n).font = font(9, color=st.MUTED, italic=True)
    ws.column_dimensions["B"].width = 52
    for col in "CDEFGIJKL":
        ws.column_dimensions[col].width = 15
    ws.column_dimensions["H"].width = 3
    ws.freeze_panes = "C7"
    ws.protection.sheet = True


def per_property(wb):
    ws = wb.create_sheet("Per Property")
    title(ws, "Per Property", "Each property for the whole tax year. Nothing to type here.")
    header(ws, 5, ["Property", "Your share", "Income (100%)", "Expenses (100%)", "Finance costs (100%)",
                   "Profit (100%)", "Your share of profit"])
    rng = lambda sheet, col: f"{sheet}!${col}${FIRST}:${col}${LAST}"  # noqa: E731
    for i in range(PROPS):
        r = 6 + i
        p = f"$B{r}"
        ws.cell(row=r, column=2, value=f'=IF(Settings!B{PROP_FIRST + i}="","",Settings!B{PROP_FIRST + i})')
        ws.cell(row=r, column=3, value=f'=IF(B{r}="","",Settings!C{PROP_FIRST + i})').number_format = "0%"
        ws.cell(row=r, column=4, value=f'=IF(B{r}="","",SUMIFS({rng("Income", "F")},{rng("Income", "C")},{p},{rng("Income", "G")},"Q*"))')
        ws.cell(row=r, column=5, value=(f'=IF(B{r}="","",SUMIFS({rng("Expenses", "F")},{rng("Expenses", "C")},{p},{rng("Expenses", "G")},"Q*")'
                                        f'-SUMIFS({rng("Expenses", "F")},{rng("Expenses", "C")},{p},{rng("Expenses", "G")},"Q*",{rng("Expenses", "E")},"{FINANCE}"))'))
        ws.cell(row=r, column=6, value=f'=IF(B{r}="","",SUMIFS({rng("Expenses", "F")},{rng("Expenses", "C")},{p},{rng("Expenses", "G")},"Q*",{rng("Expenses", "E")},"{FINANCE}"))')
        ws.cell(row=r, column=7, value=f'=IF(B{r}="","",D{r}-E{r})')
        ws.cell(row=r, column=8, value=f'=IF(B{r}="","",G{r}*C{r})')
    t = 6 + PROPS
    ws.cell(row=t, column=2, value="All properties")
    for col in "DEFGH":
        ws[f"{col}{t}"] = f"=SUM({col}6:{col}{t - 1})"
    for r in range(6, t + 1):
        for col in range(2, 9):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            if col > 3:
                c.number_format = GBP
            if r == t:
                c.font = font(10, True)
                c.fill = PatternFill("solid", fgColor=BRAND_SOFT)
    ws.cell(row=t + 2, column=2, value="Profit here is before finance costs, which are not deductible for residential lets.").font = font(9, color=st.MUTED, italic=True)
    ws.column_dimensions["B"].width = 28
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 17
    ws.protection.sheet = True


def year_end(wb):
    ws = wb.create_sheet("Year End")
    title(ws, "Year End totals", "SA105-style totals, your share. Nothing to type here.")
    ws["B3"] = '="Tax year "&Settings!C9&"   ·   final declaration due by "&TEXT(Settings!E17,"dd/mm/yyyy")'
    ws["B3"].font = font(10, True, st.MUTED)
    header(ws, 5, ["Line (as on the SA105 UK property pages)", "Your share £"])
    rows = [("Income from property (rents and other income)",
             f'=SUMIFS(Income!$I${FIRST}:$I${LAST},Income!$G${FIRST}:$G${LAST},"Q*")')]
    for c in DEDUCTIBLE:
        rows.append((c, f'=SUMIFS(Expenses!$I${FIRST}:$I${LAST},Expenses!$E${FIRST}:$E${LAST},$B{6 + len(rows)},Expenses!$G${FIRST}:$G${LAST},"Q*")'))
    for i, (label, f) in enumerate(rows):
        ws.cell(row=6 + i, column=2, value=label)
        ws.cell(row=6 + i, column=3, value=f)
    r = 6 + len(rows)
    te, np_, fin = r, r + 1, r + 2
    ws.cell(row=te, column=2, value="Total allowable expenses")
    ws.cell(row=te, column=3, value=f"=SUM(C7:C{te - 1})")
    ws.cell(row=np_, column=2, value="Net property profit (or loss if negative)")
    ws.cell(row=np_, column=3, value=f"=C6-C{te}")
    ws.cell(row=fin, column=2, value=FINANCE + " (kept apart: 20% tax reduction)")
    ws.cell(row=fin, column=3, value=f'=SUMIFS(Expenses!$I${FIRST}:$I${LAST},Expenses!$E${FIRST}:$E${LAST},"{FINANCE}",Expenses!$G${FIRST}:$G${LAST},"Q*")')
    for rr in range(6, fin + 1):
        for col in (2, 3):
            c = ws.cell(row=rr, column=col)
            c.border = BOX
            if col == 3:
                c.number_format = GBP
            if rr in (te, np_):
                c.font = font(10, True)
                c.fill = PatternFill("solid", fgColor=BRAND_SOFT)

    a = fin + 2
    header(ws, a, ["Property allowance check", ""])
    ws.cell(row=a + 1, column=2, value="Property allowance")
    ws.cell(row=a + 1, column=3, value="=Settings!C27").number_format = GBP
    ws.cell(row=a + 2, column=2, value=(f'=IF(C{te}>=C{a + 1},"Your expenses are higher than the allowance: claim expenses.",'
                                        f'"The allowance is higher than your expenses. If you use it, you cannot also deduct expenses or claim the finance cost reduction.")'))
    ws.cell(row=a + 2, column=2).font = font(10, True, BRAND)

    m = a + 4
    header(ws, m, ["Making Tax Digital check", ""])
    ws.cell(row=m + 1, column=2, value="Property income, your share (gross)")
    ws.cell(row=m + 1, column=3, value="=C6").number_format = GBP
    ws.cell(row=m + 2, column=2, value="Self-employment turnover (from Settings)")
    ws.cell(row=m + 2, column=3, value="=Settings!C7").number_format = GBP
    ws.cell(row=m + 3, column=2, value="Qualifying income")
    ws.cell(row=m + 3, column=3, value=f"=C{m + 1}+C{m + 2}").number_format = GBP
    q = f"C{m + 3}"
    ws.cell(row=m + 4, column=2, value=(f'=IF({q}>50000,"Over £50,000: Making Tax Digital applies from April 2026",'
                                        f'IF({q}>30000,"Over £30,000: Making Tax Digital applies from April 2027",'
                                        f'IF({q}>20000,"Over £20,000: Making Tax Digital applies from April 2028",'
                                        f'"£20,000 or less: not yet in Making Tax Digital")))'))
    ws.cell(row=m + 4, column=2).font = font(10, True, BRAND)
    ws.cell(row=m + 5, column=2, value="HMRC tests the tax year before. Treat this as a guide and check GOV.UK.").font = font(9, color=st.MUTED, italic=True)
    for rr in range(m + 1, m + 4):
        for col in (2, 3):
            ws.cell(row=rr, column=col).border = BOX
    notes = ["Improvements (extensions, upgrades) are capital costs, not expenses. Like-for-like replacement of furniture "
             "and appliances goes in 'Costs of replacing domestic items'.",
             "Furnished holiday lets are treated as ordinary lets from 6 April 2025.",
             "This is a record-keeping tool, not tax advice. It does not work out the tax you owe."]
    for i, n in enumerate(notes):
        ws.cell(row=m + 7 + i, column=2, value=n).font = font(9, color=st.MUTED, italic=True)
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 18
    ws.protection.sheet = True


def start(wb):
    ws = wb.active
    ws.title = "Start Here"
    title(ws, "UK Landlord Rental Tracker", "Rent and costs for up to 5 properties, totalled by Making Tax Digital quarter and SA105 line.")
    lines = [
        "SET UP IN 3 STEPS",
        "1.  Go to Settings. Type your name, the tax year (2026 for 2026-27) and up to 5 properties with your share of each.",
        "2.  Record rent on Income and costs on Expenses, choosing the property each time. Type only in the cream cells.",
        "3.  Open Quarterly Summary for each update, Per Property to compare lets, and Year End for your SA105 totals.",
        "",
        "GOOD TO KNOW",
        "•  Enter the full amount for the property. The Your share column works out your part from the % in Settings.",
        "•  Mortgage interest goes under 'Residential property finance costs'. It is kept out of profit because it gives a 20% tax reduction instead.",
        "•  The Quarter column fills itself from the date. 'Outside year' in red means the date is not in this tax year.",
        "•  Grey cells hold formulas and are protected without a password (Review > Unprotect Sheet to change them).",
        "•  Works in Excel, Google Sheets (File > Import), Apple Numbers and LibreOffice. No macros.",
        "",
        "IMPORTANT",
        "This spreadsheet keeps your records and adds them up. It does not submit anything to HMRC. Use HMRC-recognised "
        "bridging or MTD software to send quarterly updates.",
        "It is a record-keeping tool, not tax advice. Check anything you are unsure about with HMRC guidance or an accountant.",
        "Not for limited companies or overseas property.",
    ]
    for i, text in enumerate(lines):
        c = ws.cell(row=4 + i, column=2, value=text)
        c.font = font(10, True, BRAND) if text.isupper() and text else font(10, color=st.WARN if text.startswith(("This spreadsheet", "It is a", "Not for")) else st.INK)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[4 + i].height = 30 if len(text) > 110 else 18
    ws.column_dimensions["B"].width = 110


SAMPLE_PROPS = [("12 Elm Road", 1.0), ("Flat 3, Quay House", 0.5)]
SAMPLE_INCOME = [  # date, property, description, type, amount, ref
    (date(2026, 4, 28), "12 Elm Road", "April rent", "Rent", 950.00, "SO"),
    (date(2026, 5, 28), "12 Elm Road", "May rent", "Rent", 950.00, "SO"),
    (date(2026, 6, 28), "12 Elm Road", "June rent", "Rent", 950.00, "SO"),
    (date(2026, 4, 15), "Flat 3, Quay House", "Rent Apr-Jun", "Rent", 3300.00, "Agent stmt"),
    (date(2026, 7, 28), "12 Elm Road", "Rent Jul-Sep", "Rent", 2850.00, "SO"),
    (date(2026, 7, 15), "Flat 3, Quay House", "Rent Jul-Sep", "Rent", 3300.00, "Agent stmt"),
    (date(2026, 10, 28), "12 Elm Road", "Rent Oct-Dec", "Rent", 2850.00, "SO"),
    (date(2026, 10, 15), "Flat 3, Quay House", "Rent Oct-Dec", "Rent", 3300.00, "Agent stmt"),
    (date(2026, 11, 2), "12 Elm Road", "Deposit kept for damage", "Other property income", 180.00, "DPS"),
    (date(2027, 1, 28), "12 Elm Road", "Rent Jan-Mar", "Rent", 2850.00, "SO"),
    (date(2027, 1, 15), "Flat 3, Quay House", "Rent Jan-Mar", "Rent", 3300.00, "Agent stmt"),
    (date(2027, 4, 15), "Flat 3, Quay House", "Next year's rent", "Rent", 3300.00, "Agent stmt"),
]
SAMPLE_EXPENSES = [  # date, property, description, category, amount, ref
    (date(2026, 4, 30), "Flat 3, Quay House", "Letting agent fee 12%", "Legal, management and other professional fees", 396.00, "A01"),
    (date(2026, 5, 3), "12 Elm Road", "Landlord insurance", "Rent, rates, insurance and ground rents", 312.00, "R01"),
    (date(2026, 5, 20), "12 Elm Road", "Mortgage interest Apr-Jun", FINANCE, 1260.00, "M01"),
    (date(2026, 6, 11), "12 Elm Road", "Boiler service", "Property repairs and maintenance", 95.00, "R02"),
    (date(2026, 7, 31), "Flat 3, Quay House", "Letting agent fee 12%", "Legal, management and other professional fees", 396.00, "A02"),
    (date(2026, 8, 20), "Flat 3, Quay House", "Service charge and ground rent", "Rent, rates, insurance and ground rents", 840.00, "R03"),
    (date(2026, 9, 9), "12 Elm Road", "Replace washing machine (like for like)", "Costs of replacing domestic items", 329.00, "R04"),
    (date(2026, 10, 31), "Flat 3, Quay House", "Letting agent fee 12%", "Legal, management and other professional fees", 396.00, "A03"),
    (date(2026, 11, 20), "Flat 3, Quay House", "Mortgage interest Jul-Dec", FINANCE, 2100.00, "M02"),
    (date(2026, 12, 4), "12 Elm Road", "Gas safety certificate", "Other allowable property expenses", 70.00, "R05"),
    (date(2027, 1, 31), "Flat 3, Quay House", "Letting agent fee 12%", "Legal, management and other professional fees", 396.00, "A04"),
    (date(2027, 2, 14), "12 Elm Road", "Roof repair", "Property repairs and maintenance", 640.00, "R06"),
    (date(2027, 3, 20), "12 Elm Road", "Accountant", "Legal, management and other professional fees", 240.00, "R07"),
]


def build(path, example):
    wb = Workbook()
    start(wb)
    settings = build_settings(wb)
    props = f"=Settings!$B${PROP_FIRST}:$B${PROP_FIRST + PROPS - 1}"
    inc = st.register(wb, "Income", "Money in. Enter the full amount for the property.",
                      ["Date", "Property", "Description", "Type", "Amount £", "Quarter", "Ref"],
                      [12, 22, 28, 22, 13, 13, 12], "BCDEFH",
                      {"C": props, "E": f"=Settings!$I$4:$I${3 + len(INCOME_TYPES)}"}, "G", "F")
    exp = st.register(wb, "Expenses", "Money out. Choose the SA105 line from the list.",
                      ["Date", "Property", "Description", "Category", "Amount £", "Quarter", "Ref"],
                      [12, 22, 32, 44, 13, 13, 12], "BCDEFH",
                      {"C": props, "E": f"=Settings!$G$4:$G${3 + len(CATEGORIES)}"}, "G", "F")
    share_column(inc)
    share_column(exp)
    summary(wb)
    per_property(wb)
    year_end(wb)
    if example:
        settings["C4"] = "Sam Patel (example)"
        settings["C7"] = 24000
        for i, (n, s) in enumerate(SAMPLE_PROPS):
            settings.cell(row=PROP_FIRST + i, column=2, value=n)
            settings.cell(row=PROP_FIRST + i, column=3, value=s)
        settings.cell(row=PROP_FIRST + 1, column=4, value="Owned jointly with spouse")
        for ws, rows in ((inc, SAMPLE_INCOME), (exp, SAMPLE_EXPENSES)):
            for i, row in enumerate(rows):
                for j, v in enumerate(row[:5]):
                    ws.cell(row=FIRST + i, column=2 + j, value=v)
                ws.cell(row=FIRST + i, column=8, value=row[5])
    wb.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "dist")
    out.mkdir(parents=True, exist_ok=True)
    build(out / "UK-Landlord-Rental-Tracker-2026-27.xlsx", example=False)
    build(out / "UK-Landlord-Rental-Tracker-EXAMPLE.xlsx", example=True)
    print("built", out)
