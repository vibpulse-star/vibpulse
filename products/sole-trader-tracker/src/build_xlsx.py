"""Build the Sole Trader Quarterly Tracker workbooks (blank + example).

Only plain formulas (IF, DATE, SUMIFS, MIN, MAX) so the file works in Excel,
Google Sheets, Numbers and LibreOffice. No macros, tables or named ranges.
"""
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

TAX_YEAR = 2026  # tax year starting 6 April 2026 (2026-27)
ROWS = 1000      # input rows per register
FIRST = 5        # first data row on register sheets
LAST = FIRST + ROWS - 1

# Expense categories in SA103F order (self-employment full pages).
CATEGORIES = [
    "Cost of goods bought for resale or goods used",
    "Construction industry - payments to subcontractors",
    "Wages, salaries and other staff costs",
    "Car, van and travel expenses",
    "Rent, rates, power and insurance costs",
    "Repairs and maintenance of property and equipment",
    "Phone, fax, stationery and other office costs",
    "Advertising and business entertainment costs",
    "Interest on bank and other loans",
    "Bank, credit card and other financial charges",
    "Irrecoverable debts written off",
    "Accountancy, legal and other professional fees",
    "Depreciation and loss or profit on sale of assets",
    "Other business expenses",
]
CAR_ROW_INDEX = CATEGORIES.index("Car, van and travel expenses")
INCOME_TYPES = ["Turnover", "Other business income"]
VEHICLES = ["Car or van", "Motorcycle", "Bicycle"]
QUARTERS = ["Q1", "Q2", "Q3", "Q4"]

INK = "1F2A37"
MUTED = "5B6675"
BRAND = "1D5C63"       # deep teal
BRAND_SOFT = "E3EFEF"
INPUT = "FFFDF2"       # pale cream for cells the buyer types in
LINE = "D5DCE0"
WARN = "B4532A"

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
GBP = '£#,##0.00;[Red]-£#,##0.00'
DMY = "dd/mm/yyyy"


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name="Arial", size=size, bold=bold, color=color, italic=italic)


def title(ws, text, sub):
    ws["B1"] = text
    ws["B1"].font = font(16, True, BRAND)
    ws["B2"] = sub
    ws["B2"].font = font(10, color=MUTED)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2


def header(ws, row, labels, start_col=2):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=label)
        c.font = font(10, True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=BRAND)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30


def unlock(cell):
    cell.protection = Protection(locked=False)


def quarter_formula(date_ref):
    s = "Settings!"
    return (f'=IF({date_ref}="","",IF({date_ref}<{s}$C$13,"Outside year",'
            f'IF({date_ref}<={s}$D$13,"Q1",IF({date_ref}<={s}$D$14,"Q2",'
            f'IF({date_ref}<={s}$D$15,"Q3",IF({date_ref}<={s}$D$16,"Q4","Outside year"))))))')


def build_settings(wb):
    ws = wb.create_sheet("Settings")
    title(ws, "Settings", "Fill in the cream cells once. Everything else updates itself.")
    rows = [
        ("Business name", "Your business name", None),
        ("Tax year starts (year)", TAX_YEAR, "Type 2026 for the 2026-27 tax year (6 April 2026 to 5 April 2027)."),
        ("Quarter dates", "Standard (6th to 5th)", "Choose Calendar only if you elected calendar quarters with HMRC."),
        ("Accounting basis", "Cash basis", "Cash basis is the default for most sole traders. For your records only."),
    ]
    for i, (label, val, note) in enumerate(rows):
        r = 4 + i
        ws.cell(row=r, column=2, value=label).font = font(10, True)
        c = ws.cell(row=r, column=3, value=val)
        c.font = font(10)
        c.fill = PatternFill("solid", fgColor=INPUT)
        c.border = BOX
        unlock(c)
        if note:
            ws.cell(row=r, column=4, value=note).font = font(9, color=MUTED, italic=True)
    ws["B9"] = "Tax year"
    ws["B9"].font = font(10, True)
    ws["C9"] = '=C5&"-"&RIGHT(C5+1,2)'
    ws["B10"] = "Quarter start day"
    ws["B10"].font = font(10, True)
    ws["C10"] = '=IF(C6="Calendar (1st to end)",1,6)'

    header(ws, 12, ["Quarter", "Starts", "Ends", "Update due by"])
    q = [(4, 0, 7, 0, 8, 0), (7, 0, 10, 0, 11, 0), (10, 0, 1, 1, 2, 1), (1, 1, 4, 1, 5, 1)]
    for i, (sm, sy, em, ey, dm, dy) in enumerate(q):
        r = 13 + i
        ws.cell(row=r, column=2, value=QUARTERS[i]).font = font(10, True)
        ws.cell(row=r, column=3, value=f"=DATE($C$5+{sy},{sm},$C$10)")
        ws.cell(row=r, column=4, value=f"=DATE($C$5+{ey},{em},$C$10)-1")
        ws.cell(row=r, column=5, value=f"=DATE($C$5+{dy},{dm},7)")
        for col in (3, 4, 5):
            ws.cell(row=r, column=col).number_format = DMY
        for col in range(2, 6):
            ws.cell(row=r, column=col).border = BOX
    ws["B17"] = "Final declaration (year end) due by"
    ws["B17"].font = font(10, True)
    ws["E17"] = "=DATE($C$5+2,1,31)"
    ws["E17"].number_format = DMY

    header(ws, 19, ["Mileage rate (simplified expenses)", "Rate per mile", "Notes"])
    rates = [
        ("Car or van, first 10,000 business miles", 0.45, "HMRC approved rates, check each April."),
        ("Car or van, after 10,000 business miles", 0.25, ""),
        ("Motorcycle", 0.24, ""),
        ("Bicycle", 0.20, ""),
        ("Car or van tier threshold (miles)", 10000, ""),
    ]
    for i, (label, val, note) in enumerate(rates):
        r = 20 + i
        ws.cell(row=r, column=2, value=label)
        c = ws.cell(row=r, column=3, value=val)
        c.fill = PatternFill("solid", fgColor=INPUT)
        c.number_format = "#,##0" if i == 4 else "£0.00"
        unlock(c)
        ws.cell(row=r, column=4, value=note).font = font(9, color=MUTED, italic=True)
        for col in range(2, 5):
            ws.cell(row=r, column=col).border = BOX

    # Dropdown source lists, kept visible so buyers can see where choices come from.
    ws["G3"] = "Dropdown lists (edit with care)"
    ws["G3"].font = font(10, True, MUTED)
    for i, v in enumerate(CATEGORIES):
        ws.cell(row=4 + i, column=7, value=v)
    for i, v in enumerate(INCOME_TYPES):
        ws.cell(row=4 + i, column=9, value=v)
    for i, v in enumerate(["Yes", "No"]):
        ws.cell(row=4 + i, column=10, value=v)
    for i, v in enumerate(VEHICLES):
        ws.cell(row=4 + i, column=11, value=v)
    for i, v in enumerate(["Standard (6th to 5th)", "Calendar (1st to end)"]):
        ws.cell(row=4 + i, column=12, value=v)
    for i, v in enumerate(["Cash basis", "Traditional (accruals)"]):
        ws.cell(row=4 + i, column=13, value=v)
    for col in "GIJKLM":
        for r in range(4, 4 + len(CATEGORIES)):
            ws[f"{col}{r}"].font = font(9, color=MUTED)

    dv = DataValidation(type="list", formula1="=$L$4:$L$5", allow_blank=False)
    dv2 = DataValidation(type="list", formula1="=$M$4:$M$5", allow_blank=False)
    ws.add_data_validation(dv)
    ws.add_data_validation(dv2)
    dv.add("C6")
    dv2.add("C7")

    for col, w in {"B": 40, "C": 24, "D": 14, "E": 16, "F": 3, "G": 50, "H": 2, "I": 22, "J": 8, "K": 14, "L": 22, "M": 22}.items():
        ws.column_dimensions[col].width = w
    ws.protection.sheet = True
    return ws


def register(wb, name, sub, labels, widths, inputs, validations, formula_col, money_col, date_col="B"):
    ws = wb.create_sheet(name)
    title(ws, name, sub)
    header(ws, 4, labels)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("B") + i)].width = w
    ws.freeze_panes = "B5"
    for r in range(FIRST, LAST + 1):
        for col in inputs:
            c = ws[f"{col}{r}"]
            c.fill = PatternFill("solid", fgColor=INPUT)
            c.border = BOX
            unlock(c)
        ws[f"{date_col}{r}"].number_format = DMY
        if money_col:
            ws[f"{money_col}{r}"].number_format = GBP
        f = ws[f"{formula_col}{r}"]
        f.value = quarter_formula(f"{date_col}{r}")
        f.font = font(10, color=MUTED)
        f.border = BOX
    for col, src in validations.items():
        dv = DataValidation(type="list", formula1=src, allow_blank=True)
        dv.error = "Pick an option from the list."
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(f"{col}{FIRST}:{col}{LAST}")
    ws.conditional_formatting.add(
        f"{formula_col}{FIRST}:{formula_col}{LAST}",
        FormulaRule(formula=[f'${formula_col}{FIRST}="Outside year"'], font=Font(color=WARN, bold=True)))
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.sort = False
    ws.protection.autoFilter = False
    ws.auto_filter.ref = f"B4:{chr(ord('B') + len(labels) - 1)}{LAST}"
    return ws


def build_summary(wb):
    ws = wb.create_sheet("Quarterly Summary")
    title(ws, "Quarterly Summary", "Your figures for each quarterly update. Nothing to type here.")
    ws["B3"] = '="Tax year "&Settings!C9&"   ·   "&Settings!C4'
    ws["B3"].font = font(10, True, MUTED)
    header(ws, 5, ["", "Q1", "Q2", "Q3", "Q4", "Year total", "", "To end of Q1", "To end of Q2", "To end of Q3", "To end of Q4"])
    ws["B5"].value = "This quarter only"
    ws["H5"].fill = PatternFill(fill_type=None)
    ws["I4"] = "Year to date (cumulative, as HMRC's quarterly updates use)"
    ws["I4"].font = font(9, True, MUTED)
    ws["B6"] = "Update due by"
    for i in range(4):
        c = ws.cell(row=6, column=3 + i, value=f"=Settings!E{13 + i}")
        c.number_format = DMY
        c.font = font(10, True, WARN)
    ws["B6"].font = font(10, True, WARN)

    rows = []  # (label, kind, payload)
    rows.append(("INCOME", "section", None))
    for t in INCOME_TYPES:
        rows.append((t, "income", t))
    rows.append(("Total income", "total_income", None))
    rows.append(("ALLOWABLE EXPENSES", "section", None))
    for cat in CATEGORIES:
        rows.append((cat, "expense", cat))
    rows.append(("   of which simplified mileage allowance", "mileage", None))
    rows.append(("Total allowable expenses", "total_exp", None))
    rows.append(("NET PROFIT (income less allowable expenses)", "profit", None))
    rows.append(("Disallowable expenses recorded (not deducted)", "disallow", None))

    r = 7
    pos = {}
    for label, kind, payload in rows:
        pos.setdefault(kind, []).append(r)
        ws.cell(row=r, column=2, value=label)
        if kind == "section":
            ws.cell(row=r, column=2).font = font(9, True, BRAND)
            r += 1
            continue
        for qi, q in enumerate(QUARTERS):
            col = 3 + qi
            if kind == "income":
                f = f'=SUMIFS(Income!$F${FIRST}:$F${LAST},Income!$E${FIRST}:$E${LAST},"{payload}",Income!$G${FIRST}:$G${LAST},"{q}")'
            elif kind == "expense":
                f = (f'=SUMIFS(Expenses!$F${FIRST}:$F${LAST},Expenses!$E${FIRST}:$E${LAST},$B{r},'
                     f'Expenses!$H${FIRST}:$H${LAST},"{q}",Expenses!$G${FIRST}:$G${LAST},"<>No")')
                if payload == CATEGORIES[CAR_ROW_INDEX]:
                    f += f"+Mileage!$K${14 + qi}"
            elif kind == "mileage":
                f = f"=Mileage!$K${14 + qi}"
            elif kind == "disallow":
                f = (f'=SUMIFS(Expenses!$F${FIRST}:$F${LAST},Expenses!$H${FIRST}:$H${LAST},"{q}",'
                     f'Expenses!$G${FIRST}:$G${LAST},"No")')
            else:
                f = None
            ws.cell(row=r, column=col, value=f)
        r += 1

    def fill_totals(kind, formula_for_col):
        for rr in pos[kind]:
            for col in range(3, 7):
                L = chr(ord("A") + col - 1)
                ws.cell(row=rr, column=col, value=formula_for_col(L))

    inc = pos["income"]
    exp = pos["expense"]
    fill_totals("total_income", lambda L: f"=SUM({L}{inc[0]}:{L}{inc[-1]})")
    fill_totals("total_exp", lambda L: f"=SUM({L}{exp[0]}:{L}{exp[-1]})")
    ti, te = pos["total_income"][0], pos["total_exp"][0]
    fill_totals("profit", lambda L: f"={L}{ti}-{L}{te}")

    last = r - 1
    for rr in range(8, last + 1):
        if ws.cell(row=rr, column=3).value is None:
            continue
        ws.cell(row=rr, column=7, value=f"=SUM(C{rr}:F{rr})")
        for k in range(4):
            ws.cell(row=rr, column=9 + k, value=f"=SUM($C{rr}:{chr(ord('C') + k)}{rr})")
        for col in list(range(3, 8)) + list(range(9, 13)):
            c = ws.cell(row=rr, column=col)
            c.number_format = GBP
            c.border = BOX
        ws.cell(row=rr, column=2).border = BOX
    for kind in ("total_income", "total_exp", "profit"):
        for rr in pos[kind]:
            for col in range(2, 13):
                if col == 8:
                    continue
                c = ws.cell(row=rr, column=col)
                c.font = font(10, True)
                c.fill = PatternFill("solid", fgColor=BRAND_SOFT)
    for rr in pos["mileage"] + pos["disallow"]:
        for col in range(2, 13):
            ws.cell(row=rr, column=col).font = font(9, color=MUTED, italic=True)

    note = last + 2
    ws.cell(row=note, column=2, value="Send these figures with your HMRC-recognised bridging or MTD software. "
            "This spreadsheet keeps your records and totals; it does not submit anything to HMRC.").font = font(9, color=MUTED, italic=True)
    ws.cell(row=note + 1, column=2, value="If your turnover is under £90,000 you may be able to report expenses as one total "
            "(consolidated). Check with HMRC guidance or your accountant.").font = font(9, color=MUTED, italic=True)
    ws.column_dimensions["B"].width = 52
    for col in "CDEFG":
        ws.column_dimensions[col].width = 15
    ws.column_dimensions["H"].width = 3
    for col in "IJKL":
        ws.column_dimensions[col].width = 15
    ws.freeze_panes = "C7"
    ws.protection.sheet = True
    return ws, pos


def build_mileage_helper(ws):
    """Simplified mileage allowance per quarter, independent of row order.

    Car/van allowance on cumulative miles m: MIN(m,T)*r1 + MAX(m-T,0)*r2.
    Quarter allowance = allowance(cum to Q) - allowance(cum to previous Q).
    """
    ws["H11"] = "Mileage allowance worked out for you"
    ws["H11"].font = font(10, True, BRAND)
    header(ws, 13, ["Quarter", "Car/van miles", "Car/van miles YTD", "Allowance £"], start_col=8)
    s = "Settings!"
    for qi, q in enumerate(QUARTERS):
        r = 14 + qi
        ws.cell(row=r, column=8, value=q)
        ws.cell(row=r, column=9, value=f'=SUMIFS($E${FIRST}:$E${LAST},$D${FIRST}:$D${LAST},"Car or van",$F${FIRST}:$F${LAST},"{q}")')
        ws.cell(row=r, column=10, value=f"=SUM($I$14:I{r})")
        cur = f"(MIN(J{r},{s}$C$24)*{s}$C$20+MAX(J{r}-{s}$C$24,0)*{s}$C$21)"
        prev = "0" if qi == 0 else f"(MIN(J{r-1},{s}$C$24)*{s}$C$20+MAX(J{r-1}-{s}$C$24,0)*{s}$C$21)"
        other = (f'+SUMIFS($E${FIRST}:$E${LAST},$D${FIRST}:$D${LAST},"Motorcycle",$F${FIRST}:$F${LAST},"{q}")*{s}$C$22'
                 f'+SUMIFS($E${FIRST}:$E${LAST},$D${FIRST}:$D${LAST},"Bicycle",$F${FIRST}:$F${LAST},"{q}")*{s}$C$23')
        ws.cell(row=r, column=11, value=f"={cur}-{prev}{other}").number_format = GBP
        for col in range(8, 12):
            ws.cell(row=r, column=col).border = BOX
    ws.cell(row=19, column=8, value="Claim either this allowance or actual vehicle costs, not both.").font = font(9, color=MUTED, italic=True)
    for col, w in {"G": 24, "H": 12, "I": 14, "J": 16, "K": 14}.items():
        ws.column_dimensions[col].width = w


def build_year_end(wb):
    ws = wb.create_sheet("Year End")
    title(ws, "Year End totals", "SA103-style totals for your Self Assessment. Nothing to type here.")
    ws["B3"] = '="Tax year "&Settings!C9&"   ·   final declaration due by "&TEXT(Settings!E17,"dd/mm/yyyy")'
    ws["B3"].font = font(10, True, MUTED)
    header(ws, 5, ["Line (as on the SA103F form)", "Allowable £", "Disallowable £"])
    r = 6
    ws.cell(row=r, column=2, value="Turnover")
    ws.cell(row=r, column=3, value=f'=SUMIFS(Income!$F${FIRST}:$F${LAST},Income!$E${FIRST}:$E${LAST},"Turnover",Income!$G${FIRST}:$G${LAST},"Q*")')
    r += 1
    ws.cell(row=r, column=2, value="Any other business income")
    ws.cell(row=r, column=3, value=f'=SUMIFS(Income!$F${FIRST}:$F${LAST},Income!$E${FIRST}:$E${LAST},"Other business income",Income!$G${FIRST}:$G${LAST},"Q*")')
    first_exp = r + 1
    for cat in CATEGORIES:
        r += 1
        ws.cell(row=r, column=2, value=cat)
        allow = (f'=SUMIFS(Expenses!$F${FIRST}:$F${LAST},Expenses!$E${FIRST}:$E${LAST},$B{r},'
                 f'Expenses!$H${FIRST}:$H${LAST},"Q*",Expenses!$G${FIRST}:$G${LAST},"<>No")')
        if cat == CATEGORIES[CAR_ROW_INDEX]:
            allow += "+SUM(Mileage!$K$14:$K$17)"
        ws.cell(row=r, column=3, value=allow)
        ws.cell(row=r, column=4, value=(f'=SUMIFS(Expenses!$F${FIRST}:$F${LAST},Expenses!$E${FIRST}:$E${LAST},$B{r},'
                                        f'Expenses!$H${FIRST}:$H${LAST},"Q*",Expenses!$G${FIRST}:$G${LAST},"No")'))
    last_exp = r
    r += 1
    ws.cell(row=r, column=2, value="Total expenses")
    ws.cell(row=r, column=3, value=f"=SUM(C{first_exp}:C{last_exp})")
    ws.cell(row=r, column=4, value=f"=SUM(D{first_exp}:D{last_exp})")
    tot = r
    r += 1
    ws.cell(row=r, column=2, value="Net profit (or loss if negative)")
    ws.cell(row=r, column=3, value=f"=C6+C7-C{tot}")
    for rr in range(6, r + 1):
        for col in (2, 3, 4):
            c = ws.cell(row=rr, column=col)
            c.border = BOX
            if col > 2:
                c.number_format = GBP
    for rr in (tot, r):
        for col in (2, 3, 4):
            ws.cell(row=rr, column=col).font = font(10, True)
            ws.cell(row=rr, column=col).fill = PatternFill("solid", fgColor=BRAND_SOFT)
    notes = [
        "Copy these totals into your tax return or MTD software. Check line names against the current form before you file.",
        "Business entertainment is never allowable. Mark it 'No' in Expenses so it stays out of your profit.",
        "Depreciation is not allowable. Claim capital allowances instead (ask your accountant).",
        "This is a record-keeping tool, not tax advice. It does not work out the tax you owe.",
    ]
    for i, n in enumerate(notes):
        ws.cell(row=r + 2 + i, column=2, value=n).font = font(9, color=MUTED, italic=True)
    ws.column_dimensions["B"].width = 56
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.protection.sheet = True


def build_start(wb):
    ws = wb.active
    ws.title = "Start Here"
    title(ws, "Sole Trader Quarterly Tracker", "Income, expenses and mileage for UK sole traders, totalled by Making Tax Digital quarter.")
    blocks = [
        ("SET UP IN 3 STEPS", None),
        ("1.  Go to Settings and type your business name and the tax year (for 2026-27, type 2026).", None),
        ("2.  Record money in on Income, money out on Expenses, and business journeys on Mileage. Type only in the cream cells.", None),
        ("3.  Open Quarterly Summary for each update, and Year End for your Self Assessment totals.", None),
        ("", None),
        ("GOOD TO KNOW", None),
        ("•  The Quarter column fills itself from the date. 'Outside year' in red means the date is not in this tax year.", None),
        ("•  Mark an expense 'No' under Allowable if it is not a business cost (for example client entertaining).", None),
        ("•  Use one copy of this file per tax year. Make a copy for next year and change the year in Settings.", None),
        ("•  Grey cells hold formulas and are protected without a password. Review > Unprotect Sheet if you need to change them.", None),
        ("•  Works in Excel, Google Sheets (File > Import), Apple Numbers and LibreOffice. No macros.", None),
        ("", None),
        ("IMPORTANT", None),
        ("This spreadsheet keeps your records and adds them up. It does not submit anything to HMRC. To send quarterly updates "
         "under Making Tax Digital, use HMRC-recognised bridging or MTD software with these figures.", None),
        ("It is a record-keeping tool, not tax or accounting advice. Check anything you are unsure about with HMRC guidance or an accountant.", None),
        ("", None),
        ("The full guide (PDF) came with this download.", None),
    ]
    for i, (text, _) in enumerate(blocks):
        c = ws.cell(row=4 + i, column=2, value=text)
        if text.isupper() and text:
            c.font = font(10, True, BRAND)
        elif text.startswith(("This spreadsheet", "It is a record")):
            c.font = font(10, color=WARN)
        else:
            c.font = font(10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["B"].width = 110
    for i in range(len(blocks)):
        ws.row_dimensions[4 + i].height = 30 if len(blocks[i][0]) > 110 else 18


SAMPLE_INCOME = [
    (date(2026, 4, 14), "Harbour Coffee Co", "Logo refresh", "Turnover", 850.00, "INV-041"),
    (date(2026, 5, 2), "Meadow Yoga", "Website graphics", "Turnover", 1200.00, "INV-042"),
    (date(2026, 6, 20), "Pillar Estates", "Brochure layout", "Turnover", 640.00, "INV-043"),
    (date(2026, 7, 9), "Harbour Coffee Co", "Menu boards", "Turnover", 420.00, "INV-044"),
    (date(2026, 8, 18), "Northgate Gym", "Social media pack", "Turnover", 980.00, "INV-045"),
    (date(2026, 9, 30), "Envato", "Template royalties", "Other business income", 63.40, "Statement"),
    (date(2026, 10, 12), "Pillar Estates", "Signage design", "Turnover", 1500.00, "INV-046"),
    (date(2026, 12, 3), "Meadow Yoga", "Christmas campaign", "Turnover", 560.00, "INV-047"),
    (date(2027, 2, 11), "Bright Path Charity", "Annual report", "Turnover", 1750.00, "INV-048"),
    (date(2027, 3, 28), "Northgate Gym", "Rebrand", "Turnover", 2200.00, "INV-049"),
    (date(2027, 4, 9), "Harbour Coffee Co", "Next year's job", "Turnover", 300.00, "INV-050"),
]
SAMPLE_EXPENSES = [
    (date(2026, 4, 6), "Adobe", "Creative Cloud", "Phone, fax, stationery and other office costs", 56.98, "Yes", "R001"),
    (date(2026, 4, 22), "EE", "Mobile phone (business share)", "Phone, fax, stationery and other office costs", 18.00, "Yes", "R002"),
    (date(2026, 5, 15), "Hiscox", "Professional indemnity insurance", "Rent, rates, power and insurance costs", 132.50, "Yes", "R003"),
    (date(2026, 6, 4), "Trainline", "Train to client, Leeds", "Car, van and travel expenses", 46.20, "Yes", "R004"),
    (date(2026, 6, 30), "Starling", "Business account fee", "Bank, credit card and other financial charges", 2.00, "Yes", "R005"),
    (date(2026, 7, 21), "Meta", "Ads for portfolio", "Advertising and business entertainment costs", 75.00, "Yes", "R006"),
    (date(2026, 8, 2), "The Ivy", "Client lunch", "Advertising and business entertainment costs", 64.00, "No", "R007"),
    (date(2026, 9, 9), "Moo", "Business cards", "Advertising and business entertainment costs", 38.40, "Yes", "R008"),
    (date(2026, 10, 1), "Coworking Hub", "Desk rental (October)", "Rent, rates, power and insurance costs", 160.00, "Yes", "R009"),
    (date(2026, 11, 16), "Apple", "Laptop repair", "Repairs and maintenance of property and equipment", 189.00, "Yes", "R010"),
    (date(2027, 1, 7), "Smith & Co Accountants", "Year-end accounts", "Accountancy, legal and other professional fees", 300.00, "Yes", "R011"),
    (date(2027, 2, 19), "Envato", "Stock photo licence", "Cost of goods bought for resale or goods used", 29.00, "Yes", "R012"),
    (date(2027, 3, 3), "Figma", "Annual subscription", "Phone, fax, stationery and other office costs", 144.00, "Yes", "R013"),
]
SAMPLE_MILEAGE = [
    (date(2026, 4, 14), "Harbour Coffee Co meeting", "Car or van", 18, ""),
    (date(2026, 5, 2), "Meadow Yoga shoot", "Car or van", 42, ""),
    (date(2026, 8, 18), "Northgate Gym site visit", "Car or van", 26, ""),
    (date(2026, 10, 12), "Pillar Estates signage check", "Car or van", 64, ""),
    (date(2026, 12, 3), "Print shop collection", "Bicycle", 6, ""),
    (date(2027, 3, 28), "Northgate Gym launch", "Car or van", 26, ""),
]


def fill(ws, rows):
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            ws.cell(row=FIRST + i, column=2 + j, value=v)


def build(path, example):
    wb = Workbook()
    build_start(wb)
    settings = build_settings(wb)
    inc = register(wb, "Income", "Money in. One row per payment received.",
                   ["Date", "Customer or source", "Description", "Type", "Amount £", "Quarter", "Invoice or ref"],
                   [12, 26, 30, 22, 13, 13, 16], "BCDEFH", {"E": "=Settings!$I$4:$I$5"}, "G", "F")
    exp = register(wb, "Expenses", "Money out. One row per payment, with the receipt reference.",
                   ["Date", "Supplier", "Description", "Category", "Amount £", "Allowable?", "Quarter", "Receipt ref"],
                   [12, 22, 28, 46, 13, 11, 13, 13], "BCDEFGI",
                   {"E": "=Settings!$G$4:$G$17", "G": "=Settings!$J$4:$J$5"}, "H", "F")
    mil = register(wb, "Mileage", "Business journeys only. The allowance is worked out on the right.",
                   ["Date", "Journey and purpose", "Vehicle", "Business miles", "Quarter"],
                   [12, 34, 14, 14, 13], "BCDE", {"D": "=Settings!$K$4:$K$6"}, "F", None)
    for r in range(FIRST, LAST + 1):  # mileage has a notes column too
        unlock(mil[f"G{r}"])
    mil["G4"] = "Notes"
    mil["G4"].font = font(10, True, "FFFFFF")
    mil["G4"].fill = PatternFill("solid", fgColor=BRAND)
    build_mileage_helper(mil)
    build_summary(wb)
    build_year_end(wb)
    if example:
        settings["C4"] = "Rowan Ellis Design (example)"
        fill(inc, [r[:4] + (r[4], ) for r in SAMPLE_INCOME])
        for i, r in enumerate(SAMPLE_INCOME):
            inc.cell(row=FIRST + i, column=8, value=r[5])
        for i, r in enumerate(SAMPLE_EXPENSES):
            for j, v in enumerate(r[:6]):
                exp.cell(row=FIRST + i, column=2 + j, value=v)
            exp.cell(row=FIRST + i, column=9, value=r[6])
        fill(mil, [r[:4] for r in SAMPLE_MILEAGE])
    wb.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "dist")
    out.mkdir(parents=True, exist_ok=True)
    build(out / "Sole-Trader-Quarterly-Tracker-2026-27.xlsx", example=False)
    build(out / "Sole-Trader-Quarterly-Tracker-EXAMPLE.xlsx", example=True)
    print("built", out)
